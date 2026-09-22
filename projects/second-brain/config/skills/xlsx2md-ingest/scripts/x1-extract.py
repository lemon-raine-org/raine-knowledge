#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1", "pillow"]
# ///
"""X1/X2: Excel 통합문서(.xlsx/.xlsm/.xls) -> MD 추출 랩퍼 (xlsx2md-ingest).

usage: uv run x1-extract.py <file.xlsx|xlsm|xls> <out.md>
                            [--max-rows=N] [--max-cols=N]
                            [--media-dir=DIR] [--include-hidden]

경로:
  X1  .xlsx/.xlsm -> openpyxl 직행                        (기본 경로. 순수 Python)
  X2a .xls        -> LibreOffice(soffice) -> xlsx -> X1   (전 플랫폼)
  X2c .xls        -> Excel COM(PowerShell) -> xlsx -> X1  (Windows + Excel 설치 시)

.xls 우선순위: X2a -> X2c(Windows) -> 실패. openpyxl은 BIFF(.xls)를 읽지 못한다.

시트 하나가 MD 표 하나가 된다 (H2 = 시트명). 잉게스트는 wiki 지식이 목적이라 표 전체를
쏟지 않고 --max-rows/--max-cols 로 자른다 (기본 200행 x 40열). 잘린 자리에는
`<!-- truncated ... -->` 주석이 남고 stats truncated=1 이 뜬다. 원본 전체는 raw/xls/ 에
그대로 보존되므로 잘림은 손실이 아니다.

수식 셀은 캐시된 계산값을 쓴다. 캐시가 없으면(LibreOffice·스크립트가 만든 파일) 수식
원문(`=SUM(...)`)을 그대로 적고 stats uncached 로 센다 — 빈칸으로 지우지 않는다.

stdout 마지막 줄에 판정용 통계를 출력한다:
  stats cells=<비어있지 않은 셀> sheets=<렌더된 시트> tables=<MD 표> images=<n> charts=<n>
        formulas=<n> uncached=<캐시 없는 수식> merges=<병합 범위> truncated=<0|1>
        hidden=<건너뛴 숨김 시트> path=<X1|X2a|X2c> title=<제목|->
"""
import datetime
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

IS_WINDOWS = sys.platform == "win32"
IS_MACOS = sys.platform == "darwin"

DEFAULT_MAX_ROWS = 200   # wiki 잉게스트는 데이터 덤프가 아니라 의미 파악이 목적이다
DEFAULT_MAX_COLS = 40
SCAN_ROW_CAP = 20000     # 서식만 남은 셀로 max_row가 1048576까지 부풀 때의 안전판
SCAN_COL_CAP = 200
MEDIA_SUBDIR = "media"
TIMEOUT_SEC = 300        # LibreOffice 첫 실행 프롬프트·Excel COM 모달에서 무기한 정지 방지
SUBPROCESS_IO = dict(capture_output=True, text=True, encoding="utf-8", errors="replace")


def run(cmd: list[str], *, fatal: bool = True) -> bool:
    """명령 실행. fatal=True면 실패 시 종료, False면 stderr 경고 후 False 반환."""
    try:
        proc = subprocess.run(cmd, timeout=TIMEOUT_SEC, **SUBPROCESS_IO)
    except subprocess.TimeoutExpired:
        msg = f"command timed out after {TIMEOUT_SEC}s: {' '.join(cmd[:2])}"
        if fatal:
            sys.exit(msg)
        print(f"warning: {msg}", file=sys.stderr)
        return False
    if proc.returncode != 0:
        detail = (proc.stderr or proc.stdout or "").strip().splitlines()
        msg = f"command failed: {' '.join(cmd[:2])} — {detail[0] if detail else 'no output'}"
        if fatal:
            sys.exit(msg)
        print(f"warning: {msg}", file=sys.stderr)
        return False
    return True


# --- .xls 사전 변환 -------------------------------------------------------

def find_soffice() -> str | None:
    """PATH -> Windows 표준 설치 경로 -> macOS 앱 번들 순으로 LibreOffice를 찾는다.
    Windows 설치기는 PATH를 안 건드리므로 PATH만 보면 항상 못 찾는다."""
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    candidates: list[Path] = []
    if IS_WINDOWS:
        for env in ("ProgramFiles", "ProgramFiles(x86)"):
            base = os.environ.get(env)
            if base:
                candidates.append(Path(base) / "LibreOffice" / "program" / "soffice.exe")
    elif IS_MACOS:
        candidates.append(Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"))
    for c in candidates:
        if c.is_file():
            return str(c)
    return None


def xls_via_soffice(src: Path, work: Path) -> Path | None:
    """X2a: .xls -> .xlsx (LibreOffice headless). 없거나 실패하면 None (다음 폴백으로).
    -env:UserInstallation 으로 실행별 프로필을 준다 — GUI LibreOffice가 떠 있으면
    headless --convert-to 가 rc 0으로 아무것도 안 만드는 알려진 동작을 피한다."""
    soffice = find_soffice()
    if not soffice:
        return None
    profile = (work / "lo-profile").as_uri()
    ok = run(
        [soffice, f"-env:UserInstallation={profile}", "--headless",
         "--convert-to", "xlsx", "--outdir", str(work), str(src)],
        fatal=False,
    )
    produced = work / (src.stem + ".xlsx")
    if ok and produced.exists():
        return produced
    print("warning: X2a(LibreOffice) 변환이 산출물을 내지 않았다 — 다음 폴백으로", file=sys.stderr)
    return None


def xls_via_excel_com(src: Path, work: Path) -> Path | None:
    """X2c: .xls -> .xlsx (Excel COM 자동화, PowerShell). Windows + Excel 설치 시에만.
    ReadOnly로 열어 SaveAs(51 = xlOpenXMLWorkbook)로 사본을 저장한다 — 원본 무수정."""
    if not IS_WINDOWS:
        return None
    powershell = shutil.which("powershell") or shutil.which("pwsh")
    if not powershell:
        return None
    produced = work / (src.stem + ".xlsx")

    def ps_quote(p: Path) -> str:
        return "'" + str(p.resolve()).replace("'", "''") + "'"

    script = work / "excel2xlsx.ps1"
    script.write_text(
        "$ErrorActionPreference = 'Stop'\n"
        # Write-Error 는 EAP=Stop 아래서 종료 예외라 exit 2 에 못 미친다 — Console 직접 출력
        "try { $xl = New-Object -ComObject Excel.Application } "
        "catch { [Console]::Error.WriteLine('Excel COM unavailable (Excel not installed?)'); exit 2 }\n"
        "$xl.Visible = $false\n"
        "$xl.DisplayAlerts = $false\n"
        "try {\n"
        f"  $wb = $xl.Workbooks.Open({ps_quote(src)}, 0, $true)\n"
        f"  $wb.SaveAs({ps_quote(produced)}, 51)\n"
        "  $wb.Close($false)\n"
        "} finally { $xl.Quit() }\n",
        encoding="utf-8-sig",
    )
    try:
        proc = subprocess.run(
            [powershell, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
             "-File", str(script)],
            timeout=TIMEOUT_SEC, **SUBPROCESS_IO,
        )
    except subprocess.TimeoutExpired:
        sys.exit(f"command timed out after {TIMEOUT_SEC}s: Excel COM (대화상자가 떠 있는지 확인)")
    if proc.returncode == 2:
        return None  # Excel 없음 — 다음 폴백으로
    if proc.returncode != 0:
        detail = (proc.stderr or proc.stdout or "").strip().splitlines()
        sys.exit(f"command failed: Excel COM — {detail[0] if detail else 'no output'}")
    return produced if produced.exists() else None


# --- 셀 값 정규화 ---------------------------------------------------------

CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def fmt_value(v) -> str:
    """셀 값을 MD 표에 넣을 문자열로. 날짜는 ISO, 정수형 float는 정수로."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime.datetime):
        if (v.hour, v.minute, v.second, v.microsecond) == (0, 0, 0, 0):
            return v.date().isoformat()
        return v.isoformat(sep=" ", timespec="seconds")
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    if isinstance(v, datetime.timedelta):
        return str(v)
    if isinstance(v, float):
        r = round(v, 10)  # 0.1+0.2 부동소수 노이즈 제거
        if r == int(r) and abs(r) < 1e15:
            return str(int(r))
        return repr(r)
    return str(v)


def escape_cell(s: str) -> str:
    """MD 표 셀로 안전하게. 파이프는 셀을 쪼개고, 줄바꿈은 표를 깬다."""
    s = CONTROL.sub("", s)
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    s = s.replace("|", "\\|").replace("\n", "<br>")
    return s.strip()


def formula_text(raw) -> str | None:
    """수식 셀이면 수식 문자열, 아니면 None. ArrayFormula/SharedFormula도 잡는다."""
    if isinstance(raw, str):
        return raw if raw.startswith("=") else None
    text = getattr(raw, "text", None)  # openpyxl ArrayFormula
    if isinstance(text, str) and text.startswith("="):
        return text
    return None


# --- 시트 -> MD ----------------------------------------------------------

def column_letter(idx: int) -> str:
    out = ""
    while idx > 0:
        idx, rem = divmod(idx - 1, 26)
        out = chr(65 + rem) + out
    return out


class Counters:
    def __init__(self) -> None:
        self.cells = 0
        self.formulas = 0
        self.uncached = 0
        self.merges = 0
        self.images = 0
        self.charts = 0
        self.truncated = False


def sheet_grid(ws_v, ws_f, counters: Counters) -> list[list[str]]:
    """시트를 문자열 격자로. 사용 범위를 넘는 빈 행·열은 잘라낸다."""
    max_row = min(ws_f.max_row or 0, SCAN_ROW_CAP)
    max_col = min(ws_f.max_column or 0, SCAN_COL_CAP)
    if max_row < 1 or max_col < 1:
        return []
    grid: list[list[str]] = []
    for r in range(1, max_row + 1):
        row: list[str] = []
        for c in range(1, max_col + 1):
            raw_f = ws_f.cell(row=r, column=c).value
            raw_v = ws_v.cell(row=r, column=c).value
            fml = formula_text(raw_f)
            if fml is not None:
                counters.formulas += 1
                if raw_v is None:
                    counters.uncached += 1
                    text = fml  # 계산 캐시가 없다 — 수식 원문을 남긴다
                else:
                    text = fmt_value(raw_v)
            else:
                text = fmt_value(raw_v if raw_v is not None else raw_f)
            text = escape_cell(text)
            if text:
                counters.cells += 1
            row.append(text)
        grid.append(row)

    grid = [row for row in grid if any(row)]  # 스페이서 빈 행은 버린다 (행 번호가 원본과 어긋난다)
    while grid and all(not row[-1] for row in grid):
        for row in grid:
            row.pop()
        if not grid[0]:
            return []
    return grid


def grid_to_table(grid: list[list[str]], max_rows: int, max_cols: int,
                  counters: Counters) -> list[str]:
    """격자를 GFM 표로. 첫 행이 헤더이며, 비어 있으면 열 문자(A, B, ...)를 헤더로 쓴다."""
    total_rows, total_cols = len(grid), len(grid[0])
    rows = grid[:max_rows]
    cols = min(total_cols, max_cols)
    rows = [row[:cols] for row in rows]

    header = rows[0]
    if not any(header):
        header = [column_letter(i + 1) for i in range(cols)]
        body = rows[1:]
    else:
        body = rows[1:]
    header = [h or column_letter(i + 1) for i, h in enumerate(header)]

    out = ["| " + " | ".join(header) + " |",
           "|" + "|".join([" --- "] * cols) + "|"]
    out += ["| " + " | ".join(row) + " |" for row in body]

    notes = []
    if total_rows > max_rows:
        notes.append(f"rows {max_rows + 1}-{total_rows} of {total_rows}")
    if total_cols > max_cols:
        notes.append(f"cols {max_cols + 1}-{total_cols} of {total_cols}")
    if notes:
        counters.truncated = True
        out.append("")
        out.append(f"<!-- truncated: {', '.join(notes)} — 원본은 raw/xls/ 에 온전히 보존 -->")
    return out


ABS_TARGET = re.compile(r"!\[[^\]]*\]\(\s*<?([^)\s>]+)")
ABS_PATH = re.compile(r"^(?:[A-Za-z]:[\\/]|/|file:)")


def leaked_absolute(md: str) -> bool:
    """이미지 링크 대상이 절대경로면 True. 참조는 항상 media/ 상대경로여야 한다."""
    return any(ABS_PATH.match(m.group(1)) for m in ABS_TARGET.finditer(md))


def extract_images(ws, media_dir: Path | None, counters: Counters, stem: str) -> list[str]:
    """시트에 박힌 이미지를 media/ 로 빼고 MD 참조를 돌려준다. 실패해도 변환은 계속한다.
    openpyxl은 이미지 접근에 공개 API가 없어 _images 를 쓴다 (3.1 기준 안정)."""
    images = list(getattr(ws, "_images", []) or [])
    counters.images += len(images)
    if not images or media_dir is None:
        return []
    refs: list[str] = []
    target = media_dir / MEDIA_SUBDIR
    target.mkdir(parents=True, exist_ok=True)
    for i, img in enumerate(images, 1):
        try:
            data = img._data()
            ext = (getattr(img, "format", None) or "png").lower()
            name = f"{stem}-{i}.{ext}"
            (target / name).write_bytes(data)
            refs.append(f"![{stem} image {i}]({MEDIA_SUBDIR}/{name})")
        except Exception as exc:  # 이미지 하나 때문에 변환 전체를 버리지 않는다
            print(f"warning: 이미지 {i} 추출 실패 ({exc}) — 계수만 남긴다", file=sys.stderr)
    return refs


def workbook_to_md(wb_v, wb_f, title: str, opts) -> tuple[str, Counters, int, int, int]:
    counters = Counters()
    lines = [f"# {title}", ""]
    rendered = 0
    tables = 0
    hidden = 0
    for name in wb_f.sheetnames:
        ws_f, ws_v = wb_f[name], wb_v[name]
        if ws_f.sheet_state != "visible" and not opts["include_hidden"]:
            hidden += 1
            continue
        grid = sheet_grid(ws_v, ws_f, counters)
        merges = len(getattr(ws_f, "merged_cells", None).ranges) if getattr(ws_f, "merged_cells", None) else 0
        counters.merges += merges
        charts = len(getattr(ws_f, "_charts", []) or [])
        counters.charts += charts
        stem = re.sub(r"[^\w.-]+", "-", name).strip("-") or f"sheet{rendered + 1}"
        img_refs = extract_images(ws_f, opts["media_dir"], counters, stem)

        if not grid and not charts and not img_refs:
            continue
        rendered += 1
        lines.append(f"## {name}")
        lines.append("")
        if ws_f.sheet_state != "visible":
            lines.append(f"> 숨김 시트 (`sheet_state={ws_f.sheet_state}`).")
            lines.append("")
        if merges:
            lines.append(f"> 병합 셀 {merges}개 — MD 표에는 병합이 없어 좌상단 값만 남고 "
                         "나머지 칸은 빈칸이다.")
            lines.append("")
        if charts:
            lines.append(f"> 차트 {charts}개 — 텍스트로 옮기지 못한다. 원본 확인 필요.")
            lines.append("")
        if grid:
            tables += 1
            lines += grid_to_table(grid, opts["max_rows"], opts["max_cols"], counters)
            lines.append("")
        for ref in img_refs:
            lines.append(ref)
            lines.append("")
    return "\n".join(lines).rstrip() + "\n", counters, rendered, tables, hidden


# --- 진입점 --------------------------------------------------------------

FLAGS = {"--max-rows", "--max-cols", "--media-dir"}


def parse_args(argv: list[str]) -> tuple[Path, Path, dict]:
    positional: list[str] = []
    opts = {"max_rows": DEFAULT_MAX_ROWS, "max_cols": DEFAULT_MAX_COLS,
            "media_dir": None, "include_hidden": False}
    for a in argv:
        if a == "--include-hidden":
            opts["include_hidden"] = True
        elif a.startswith("--"):
            if "=" not in a:
                sys.exit(f"flag needs a value: {a} (use {'/'.join(sorted(FLAGS))}=VALUE)")
            key, val = a.split("=", 1)
            if key not in FLAGS:
                sys.exit(f"unknown flag: {key}")
            if key == "--media-dir":
                opts["media_dir"] = Path(val).expanduser().resolve()
            else:
                if not val.isdigit() or int(val) < 1:
                    sys.exit(f"{key} must be a positive integer: {val}")
                opts["max_rows" if key == "--max-rows" else "max_cols"] = int(val)
        else:
            positional.append(a)
    if len(positional) != 2:
        sys.exit(__doc__)
    return Path(positional[0]).expanduser(), Path(positional[1]).expanduser(), opts


def main() -> None:
    src, out, opts = parse_args(sys.argv[1:])
    if not src.is_file():
        sys.exit(f"not a file: {src}")

    suffix = src.suffix.lower()
    if suffix not in (".xlsx", ".xlsm", ".xls"):
        sys.exit(f"unsupported extension: {suffix} (.xlsx/.xlsm/.xls만 지원. "
                 "csv는 이미 텍스트라 변환이 필요 없다)")

    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl이 없다. `uv run`으로 실행하거나 `pip install openpyxl`")

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        if suffix == ".xls":
            converted = xls_via_soffice(src, work)
            path_used = "X2a"
            if converted is None:
                converted = xls_via_excel_com(src, work)
                path_used = "X2c"
            if converted is None:
                hint = ("Windows: Excel 설치(X2c) 또는 LibreOffice 설치(X2a)."
                        if IS_WINDOWS else "LibreOffice(soffice) 설치 후 재시도.")
                sys.exit(f".xls는 openpyxl이 직접 못 읽는다(BIFF 포맷). {hint}")
            target = converted
        else:
            path_used = "X1"
            target = src

        try:
            wb_v = load_workbook(target, data_only=True)
            wb_f = load_workbook(target, data_only=False)
        except Exception as exc:
            sys.exit(f"통합문서를 열 수 없다 ({type(exc).__name__}: {exc}) — "
                     "암호 걸린 파일이거나 손상된 파일일 수 있다")

        title = (getattr(wb_f.properties, "title", None) or "").strip() or src.stem
        md, counters, sheets, tables, hidden = workbook_to_md(wb_v, wb_f, title, opts)
        wb_v.close()
        wb_f.close()

    if not md.strip() or sheets == 0:
        sys.exit("empty conversion result — 렌더할 시트가 없다 "
                 "(전부 비었거나 전부 숨김. --include-hidden 확인)")

    if leaked_absolute(md):
        sys.exit("absolute path leaked into MD image links — 산출물을 쓰지 않는다")

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(md, encoding="utf-8")

    if counters.uncached:
        print(f"warning: 계산 캐시가 없는 수식 {counters.uncached}개 — 수식 원문을 그대로 "
              "남겼다. Excel에서 한 번 열고 저장하면 값으로 바뀐다.", file=sys.stderr)
    print(f"wrote {len(md)} chars -> {out}")
    print(
        f"stats cells={counters.cells} sheets={sheets} tables={tables} "
        f"images={counters.images} charts={counters.charts} "
        f"formulas={counters.formulas} uncached={counters.uncached} "
        f"merges={counters.merges} truncated={int(counters.truncated)} "
        f"hidden={hidden} path={path_used} title={title}"
    )


if __name__ == "__main__":
    main()
